"""Sinh data.js (cho trang xem) và bảng tính .xlsx từ data.json.
Chạy lại mỗi khi sửa data.json:  python3 build.py
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
D = json.loads((ROOT / "data.json").read_text(encoding="utf-8"))

(ROOT / "data.js").write_text(
    "window.POSE_DATA = " + json.dumps(D, ensure_ascii=False, indent=1) + ";\n",
    encoding="utf-8",
)

HEAD_FILL = PatternFill("solid", fgColor="1F2430")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, cols, rows):
    ws = wb.create_sheet(title)
    ws.append([c for c, _ in cols])
    for r in rows:
        ws.append(r)
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        ws.cell(1, i).fill = HEAD_FILL
        ws.cell(1, i).font = HEAD_FONT
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = "C2"
    return ws


def body(p, who, key):
    b = (p.get("co_the") or {}).get(who)
    return b[key] if b else ""


wb = Workbook()
wb.remove(wb.active)

PEOPLE = list(dict.fromkeys(k for p in D["poses"] for k in (p.get("co_the") or {})))
cols = [("ID", 7), ("Tên pose", 26), ("Chủ đề", 18), ("Bối cảnh", 16), ("Đối tượng", 11), ("Số người", 10),
        ("Trang phục hợp", 22), ("Thân mật (1-5)", 10), ("Nhịp", 12), ("Độ khó", 9), ("Tư thế nền", 11)]
for who in PEOPLE:
    cols += [(f"{who}: tay", 26), (f"{who}: chân", 22), (f"{who}: hướng mặt", 20), (f"{who}: ánh mắt", 18)]
cols += [("Đội hình", 26), ("Vị trí trong đội hình", 36), ("Quy tắc chung", 36),
         ("Tương tác", 14), ("Khung hình", 18), ("Góc máy", 18), ("Tiêu cự", 18),
         ("ĐIỂM NEO", 30), ("Lỗi hay gặp", 40),
         ("Câu Đặt", 40), ("Câu DẪN (hiện to)", 44), ("Câu Bắt", 40), ("Chỉnh trang phục / Kiểm tra", 34),
         ("Pose kế tiếp (đổi gì)", 26), ("Ảnh (Pinterest 736px)", 30), ("Ảnh thêm cùng mood", 40), ("Nguồn", 30)]
rows = []
for p in D["poses"]:
    dh = p.get("doi_hinh") or {}
    r = [p["id"], p["ten"], ", ".join(p["chu_de"]), ", ".join(p["boi_canh"]), p.get("doi_tuong", ""), p["so_nguoi"],
         ", ".join(p["trang_phuc"]), p["than_mat"] if p["than_mat"] is not None else "", p["nhip"], p["do_kho"], p["tu_the"]]
    for who in PEOPLE:
        r += [body(p, who, k) for k in ("tay", "chan", "mat", "anh_mat")]
    r += [dh.get("kieu", ""), dh.get("vi_tri", ""), dh.get("quy_tac", ""),
          p["tuong_tac"], p["khung"], p["goc"], p["tieu_cu"], p["diem_neo"], "\n".join(p["loi"]),
          p["dat"], p["dan"], p["bat"], p["chinh"],
          "\n".join(f'{n["id"]} ({n["doi"]})' for n in p["ke_tiep"]), p["anh"], "\n".join(p.get("anh_them", [])), p["nguon"]]
    rows.append(r)
sheet(wb, "Pose", cols, rows)

rows = []
for c in D["chains"]:
    for i, s in enumerate(c["steps"], 1):
        rows.append([c["id"], c["ten"], f'{c["chu_de"]} · {c["boi_canh"]} · {c["kieu"]}', i,
                     s["doi"], s.get("pose", ""), s["dang"], s["cau"], s["than_mat"]])
sheet(wb, "Chuỗi", [("Chuỗi", 7), ("Tên chuỗi", 24), ("Loại", 34), ("Bước", 7), ("Đổi trục", 12),
                    ("Pose", 8), ("Dáng", 36), ("Câu nối", 56), ("Thân mật", 10)], rows)

rows = [[g["nhom"], i, c] for g in D["lines"] for i, c in enumerate(g["cau"], 1)]
sheet(wb, "Câu dùng chung", [("Nhóm", 20), ("STT", 6), ("Câu", 80)], rows)

vocab = [
    ("Tư thế nền", "đứng · đi · ngồi · dựa · nằm · chạy · xoay người · cõng"),
    ("Tay (tự thân)", "buông tự nhiên · đút túi · chạm tóc/mặt · chỉnh cổ áo, cúc vest · cầm hoa · cầm đuôi váy · khoanh nhẹ"),
    ("Tay (lên người kia)", "đặt vai · ôm eo · đặt ngực · nâng má · nắm tay · khoác tay · ôm từ sau"),
    ("Chân", "dồn trọng tâm 1 chân · chân trước chân sau · bắt chéo · đang bước · kiễng chân · khuỵu gối nhẹ"),
    ("Hướng mặt", "chính diện · 3/4 · nghiêng hẳn · cúi nhẹ · ngẩng · nghiêng đầu"),
    ("Ánh mắt", "nhìn máy · nhìn nhau · nhìn xuống · nhìn xa · nhắm mắt · nhìn vào vật"),
    ("Tương tác", "không chạm · chạm nhẹ · nắm · khoác tay · ôm · ôm từ sau · trán chạm · mũi chạm · hôn · bế · cõng · đùa nghịch"),
    ("Đội hình nhóm", "hàng ngang · một hàng vòng cung · hai tầng · so le · tam giác"),
    ("Mức thân mật", "1 không chạm · 2 chạm nhẹ · 3 gần gũi · 4 thân mật · 5 rất thân mật"),
    ("Nhịp", "Khởi động · Xây dựng · Cao trào · Thả lỏng"),
    ("Biến xưng hô", "[nam] · [nữ] · [cặp]"),
]
sheet(wb, "Từ vựng", [("Trục", 22), ("Giá trị cho phép", 110)], vocab)

out = ROOT / "pose-coach-giai-doan-1.xlsx"
wb.save(out)
print("OK:", out.name, "+ data.js")
